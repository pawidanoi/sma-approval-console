# -*- coding: utf-8 -*-
"""อ่านไฟล์ Excel 3 ไฟล์ (สาขา, ทีมกิจ, ทีม setup) แล้วแปลงเป็น JSON สะอาดพร้อม import เข้า Supabase
ไม่แก้ไฟล์ต้นฉบับ อ่านอย่างเดียว"""
import json
import re
import unicodedata
from openpyxl import load_workbook

BRANCH_FILE = r"C:\Users\pawida.pro\Downloads\สาขาอัพเดต.xlsx"
ACTIVITY_FILE = r"C:\Users\pawida.pro\Downloads\รายชื่อทีมกิจ.xlsx"
SETUP_FILE = r"C:\Users\pawida.pro\Downloads\ทีม set up.xlsx"
OUT_DIR = r"C:\Users\pawida.pro\projects\sma-approval-console\data"

ACTIVE_STATUSES = {"Active store", "Set up ร้าน", "ก่อสร้างแล้ว รอเปิดร้าน", "ยังไม่ลงเสาเข็ม"}


def norm(s):
    if s is None:
        return None
    s = unicodedata.normalize("NFKC", str(s)).strip()
    return s if s else None


def parse_coords(raw):
    """คืน (lat, lng) หรือ (None, None) - รองรับ 3 รูปแบบที่เจอในไฟล์จริง"""
    if raw is None:
        return None, None
    s = str(raw).strip()
    if not s:
        return None, None
    # DMS: 13°45'36.1"N 99°44'25.2"E
    dms = re.findall(r"(\d+)\D+(\d+)\D+([\d.]+)\D*([NSEW])", s)
    if len(dms) == 2:
        def to_dec(deg, minute, sec, hemi):
            val = float(deg) + float(minute) / 60 + float(sec) / 3600
            return -val if hemi in ("S", "W") else val
        lat = next((to_dec(*d) for d in dms if d[3] in ("N", "S")), None)
        lng = next((to_dec(*d) for d in dms if d[3] in ("E", "W")), None)
        if lat is not None and lng is not None:
            return round(lat, 7), round(lng, 7)
    # plain "lat,long" or "lat, long"
    parts = [p.strip() for p in s.split(",")]
    if len(parts) == 2:
        try:
            return round(float(parts[0]), 7), round(float(parts[1]), 7)
        except ValueError:
            pass
    return None, None


# ---------- 1) branches ----------
wb = load_workbook(BRANCH_FILE, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(min_row=2, values_only=True))
branches = []
new_seq = 0
for r in rows:
    code_th, code_cjx, name_th, name_en, tambon, amphoe, province, address, gmap, latlng, warehouse, *_rest = r[:11]
    status = r[18] if len(r) > 18 else None
    name = norm(name_th) or norm(name_en)
    if not name:
        continue
    code = norm(code_cjx) if norm(code_cjx) and norm(code_cjx) != "-" else None
    if not code:
        new_seq += 1
        code = f"NEW-{new_seq:03d}"
    lat, lng = parse_coords(latlng)
    branches.append({
        "code": code,
        "name": name,
        "district": norm(amphoe),
        "province": norm(province),
        "address": norm(address),
        "lat": lat,
        "lng": lng,
        "status": norm(status),
        "is_active": norm(status) in ACTIVE_STATUSES,
    })
print(f"branches: {len(branches)} total, {sum(1 for b in branches if b['is_active'])} active")

# ---------- 2) activity team roster + zone muster points ----------
wb2 = load_workbook(ACTIVITY_FILE, data_only=True)
ws_roster = wb2["รายชื่อทีมกิจ"]
act_rows = list(ws_roster.iter_rows(min_row=2, values_only=True))
activity_staff = []
for r in act_rows:
    team, code, name, nickname_gender, gender, position, province, phone, address, home_latlng = r[:10]
    code = norm(code)
    if not code:
        continue
    lat, lng = parse_coords(home_latlng)
    activity_staff.append({
        "code": code, "team_code": norm(team), "name": norm(name),
        "nickname": norm(nickname_gender), "gender": "M" if norm(gender) == "ชาย" else ("F" if norm(gender) == "หญิง" else None),
        "role": norm(position), "province": norm(province), "phone": norm(phone),
        "home_lat": lat, "home_lng": lng, "category": "activity",
    })

ws_zone = wb2["จุดรวมพลทีมกิจ"]
zone_rows = list(ws_zone.iter_rows(min_row=2, values_only=True))
zones_by_province = {}  # province (base, no trailing number) -> list of {zone_label, name, lat, lng}
for r in zone_rows:
    zone, muster_name, latlng = r[:3]
    zone = norm(zone)
    if not zone:
        continue
    lat, lng = parse_coords(latlng)
    # "สมุทรปราการ 1" -> base province "สมุทรปราการ"
    m = re.match(r"^(.*?)\s*\d*$", zone)
    base_province = m.group(1).strip() if m else zone
    zones_by_province.setdefault(base_province, []).append({
        "zone_label": zone, "muster_name": muster_name, "lat": lat, "lng": lng,
    })

# resolve each activity team_code -> its province via mode of members' province
from collections import Counter
team_province = {}
for s in activity_staff:
    if s["team_code"] and s["province"]:
        team_province.setdefault(s["team_code"], Counter())[s["province"]] += 1
team_zone_resolved = {t: c.most_common(1)[0][0] for t, c in team_province.items()}

activity_muster_points = []
for team_code, province in team_zone_resolved.items():
    for z in zones_by_province.get(province, []):
        activity_muster_points.append({
            "category": "activity", "team_code": team_code, "zone_label": z["zone_label"],
            "muster_name": z["muster_name"], "lat": z["lat"], "lng": z["lng"],
        })
print(f"activity staff: {len(activity_staff)}, teams resolved to zone: {len(team_zone_resolved)}, muster rows: {len(activity_muster_points)}")

# ---------- 3) setup team roster + per-team muster points ----------
wb3 = load_workbook(SETUP_FILE, data_only=True)
ws_setup_roster = wb3["ชื่อ set up"]
setup_rows = list(ws_setup_roster.iter_rows(min_row=2, values_only=True))


def normalize_team_key(s):
    """'ราชบุรี ทีม 1' หรือ 'ราชบุรี 1' หรือ 'นครปฐม1' -> ('ราชบุรี', '1')"""
    if not s:
        return None
    s = norm(s)
    s = s.replace("ทีม", " ")
    m = re.match(r"^(.*?)\s*(\d+)\s*$", s)
    if m:
        return (m.group(1).strip(), m.group(2))
    return (s.strip(), None)


setup_staff = []
for r in setup_rows:
    team, code, name, nickname, gender, position = r[:6]
    code = norm(code)
    if not code:
        continue
    setup_staff.append({
        "code": code, "team_code": norm(team), "name": norm(name), "nickname": norm(nickname),
        "gender": "M" if norm(gender) == "ช" else ("F" if norm(gender) == "ญ" else None),
        "role": norm(position), "province": None, "phone": None,
        "home_lat": None, "home_lng": None, "category": "setup",
    })

ws_setup_zone = wb3["จุดรวมพล set up"]
setup_zone_rows = list(ws_setup_zone.iter_rows(min_row=2, values_only=True))
setup_muster_by_key = {}
for r in setup_zone_rows:
    team, muster_name, latlng = r[:3]
    key = normalize_team_key(team)
    if not key:
        continue
    lat, lng = parse_coords(latlng)
    setup_muster_by_key[key] = {"muster_name": muster_name, "lat": lat, "lng": lng}

setup_muster_points = []
seen_teams = set()
for s in setup_staff:
    if s["team_code"] in seen_teams:
        continue
    seen_teams.add(s["team_code"])
    key = normalize_team_key(s["team_code"])
    m = setup_muster_by_key.get(key)
    if m:
        setup_muster_points.append({
            "category": "setup", "team_code": s["team_code"], "zone_label": s["team_code"],
            "muster_name": m["muster_name"], "lat": m["lat"], "lng": m["lng"],
        })
matched = len(setup_muster_points)
print(f"setup staff: {len(setup_staff)}, distinct teams: {len(seen_teams)}, muster matched: {matched}")
unmatched = [t for t in seen_teams if normalize_team_key(t) not in setup_muster_by_key]
if unmatched:
    print("setup teams WITHOUT a matched muster point:", unmatched)

# ---------- write output ----------
import os
os.makedirs(OUT_DIR, exist_ok=True)
with open(os.path.join(OUT_DIR, "branches_import.json"), "w", encoding="utf-8") as f:
    json.dump(branches, f, ensure_ascii=False, indent=1)
with open(os.path.join(OUT_DIR, "staff_import.json"), "w", encoding="utf-8") as f:
    json.dump(activity_staff + setup_staff, f, ensure_ascii=False, indent=1)
with open(os.path.join(OUT_DIR, "muster_points_import.json"), "w", encoding="utf-8") as f:
    json.dump(activity_muster_points + setup_muster_points, f, ensure_ascii=False, indent=1)
print("เขียนไฟล์ JSON เสร็จแล้วที่", OUT_DIR)
